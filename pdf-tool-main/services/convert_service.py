# -*- coding: utf-8 -*-
"""PDF 转 Word/Excel/PPT/TXT/图片/HTML/EPUB 服务。"""

from __future__ import annotations

import base64
import hashlib
import html
import io
import os
import zipfile
from collections.abc import Callable

from utils.errors import AppError
from utils.files import ensure_dir

ProgressFn = Callable[[str], None]
_NOOP: ProgressFn = lambda msg: None


def _open_pdf(src: str):
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf
    if not os.path.isfile(src):
        raise AppError(f"文件不存在：{src}")
    try:
        doc = pymupdf.open(src)
    except Exception as exc:
        raise AppError(f"无法打开 PDF：{os.path.basename(src)}（{exc}）") from exc
    if doc.needs_pass:
        doc.close()
        raise AppError(f"PDF“{os.path.basename(src)}”已加密，请先移除密码")
    return doc


def _page_lines(page) -> list[str]:
    text = (page.get_text("text") or "").replace("\r", "\n")
    return [line.strip() for line in text.splitlines() if line.strip()]


def _page_tables(page):
    try:
        finder = page.find_tables()
        return [table.extract() for table in getattr(finder, "tables", []) or []
                if table.extract()]
    except Exception:
        return []


def _stem(path: str) -> str:
    return os.path.splitext(os.path.basename(path))[0]


def pdf_to_txt(src: str, out: str, progress: ProgressFn = _NOOP) -> None:
    doc = _open_pdf(src)
    try:
        parts = []
        for i, page in enumerate(doc, start=1):
            parts.append(f"======== 第 {i} 页 ========\n\n"
                         f"{(page.get_text('text') or '').strip()}")
            if i % 10 == 0:
                progress(f"已提取 {i}/{len(doc)} 页")
        with open(out, "w", encoding="utf-8") as fh:
            fh.write("\n\n".join(parts))
    finally:
        doc.close()


def pdf_to_word(src: str, out: str, progress: ProgressFn = _NOOP) -> None:
    Document = _docx()
    doc = _open_pdf(src)
    try:
        word = Document()
        for i, page in enumerate(doc, start=1):
            if i > 1:
                word.add_page_break()
            word.add_heading(f"第 {i} 页", level=2)
            for line in _page_lines(page):
                word.add_paragraph(line)
            for rows in _page_tables(page):
                if rows:
                    table = word.add_table(rows=len(rows), cols=len(rows[0]))
                    table.style = "Table Grid"
                    for r, row in enumerate(rows):
                        for c, value in enumerate(row):
                            table.cell(r, c).text = (
                                "" if value is None else str(value))
            progress(f"已写入第 {i}/{len(doc)} 页")
        word.save(out)
    finally:
        doc.close()


def pdf_to_excel(src: str, out: str, progress: ProgressFn = _NOOP) -> None:
    openpyxl = _openpyxl()
    doc = _open_pdf(src)
    try:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for i, page in enumerate(doc, start=1):
            ws = wb.create_sheet(f"第{i}页")
            for rows in _page_tables(page):
                for row in rows:
                    ws.append(["" if v is None else str(v) for v in row])
            for line in _page_lines(page):
                ws.append([line])
            progress(f"已写入第 {i}/{len(doc)} 页")
        wb.save(out)
    finally:
        doc.close()


def pdf_to_ppt(src: str, out: str, progress: ProgressFn = _NOOP) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise AppError("缺少 python-pptx，请运行 python bootstrap.py") from exc
    import pymupdf
    doc = _open_pdf(src)
    try:
        prs = Presentation()
        blank = prs.slide_layouts[6]
        slide_w, slide_h = prs.slide_width, prs.slide_height
        margin = int(slide_w * 0.02)
        for i, page in enumerate(doc, start=1):
            pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            stream = io.BytesIO(pix.tobytes("png"))
            slide = prs.slides.add_slide(blank)
            ratio = min((slide_w - margin * 2) / max(1, pix.width),
                        (slide_h - margin * 2) / max(1, pix.height))
            w = int(pix.width * ratio)
            h = int(pix.height * ratio)
            slide.shapes.add_picture(
                stream, int((slide_w - w) / 2), int((slide_h - h) / 2),
                width=w, height=h)
            progress(f"已生成第 {i}/{len(doc)} 页幻灯片")
        prs.save(out)
    finally:
        doc.close()


def pdf_to_html(src: str, out: str, progress: ProgressFn = _NOOP) -> None:
    doc = _open_pdf(src)
    try:
        title = html.escape(_stem(src))
        sections = []
        used = set()
        for i, page in enumerate(doc, start=1):
            text = html.escape(page.get_text("text") or "").replace("\n", "<br>")
            imgs = []
            for info in page.get_images(full=True):
                xref = info[0]
                if xref in used:
                    continue
                used.add(xref)
                try:
                    base = doc.extract_image(xref)
                    ext = base["ext"].lower()
                    mime = {"jpg": "image/jpeg", "jpeg": "image/jpeg",
                            "png": "image/png", "gif": "image/gif",
                            "webp": "image/webp", "bmp": "image/bmp",
                            "tif": "image/tiff", "tiff": "image/tiff"}.get(
                                ext, "image/png")
                    b64 = base64.b64encode(base["image"]).decode("ascii")
                    imgs.append(
                        f'<img src="data:{mime};base64,{b64}" '
                        f'alt="页面图片">')
                except Exception:
                    continue
            sections.append(
                f'<section><h2>第 {i} 页</h2><div>{text or "（无文本）"}</div>'
                f"{''.join(imgs)}</section>")
            if i % 5 == 0:
                progress(f"已处理 {i}/{len(doc)} 页")
        css = ("body{font-family:'Microsoft YaHei',sans-serif;line-height:1.7;"
               "margin:24px}section{page-break-after:always;"
               "border-bottom:1px solid #ccc}img{max-width:100%}")
        with open(out, "w", encoding="utf-8") as fh:
            fh.write(f'<!DOCTYPE html><html lang="zh-CN"><head>'
                     f'<meta charset="utf-8"><title>{title}</title>'
                     f'<style>{css}</style></head><body><h1>{title}</h1>'
                     f"{''.join(sections)}</body></html>")
    finally:
        doc.close()


def pdf_to_epub(src: str, out: str, progress: ProgressFn = _NOOP) -> None:
    doc = _open_pdf(src)
    try:
        title = _stem(src)
        chapters = []
        assets = {}
        counter = 0
        for i, page in enumerate(doc, start=1):
            text = html.escape(page.get_text("text") or "").replace("\n", "<br>")
            page_images = []
            for info in page.get_images(full=True):
                xref = info[0]
                if xref in assets:
                    page_images.append(f'<img src="{assets[xref][0]}" alt="">')
                    continue
                try:
                    base = doc.extract_image(xref)
                except Exception:
                    continue
                ext = base.get("ext", "png")
                counter += 1
                rel = f"images/img{counter}.{ext}"
                assets[xref] = (rel, base["image"])
                page_images.append(f'<img src="{rel}" alt="">')
            chapters.append(
                f'<html xmlns="http://www.w3.org/1999/xhtml"><head>'
                f'<title>第 {i} 页</title></head><body>'
                f'<h2>第 {i} 页</h2><div>{text or "（无文本）"}</div>'
                f"{''.join(page_images)}</body></html>")
            if i % 5 == 0:
                progress(f"已生成 {i}/{len(doc)} 个章节")

        manifest = ['<item id="nav" href="nav.xhtml" '
                    'media-type="application/xhtml+xml" properties="nav"/>']
        spine = []
        for i in range(1, len(chapters) + 1):
            manifest.append(
                f'<item id="ch{i}" href="chapter-{i:04d}.xhtml" '
                'media-type="application/xhtml+xml"/>')
            spine.append(f'<itemref idref="ch{i}"/>')
        for index, (rel, data) in enumerate(assets.values(), start=1):
            manifest.append(
                f'<item id="img{index}" href="{rel}" '
                'media-type="image/png"/>')
        uid = hashlib.sha1(title.encode("utf-8")).hexdigest()[:12]
        opf = (
            '<?xml version="1.0" encoding="utf-8"?>'
            '<package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
            '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
            f'<dc:title>{html.escape(title)}</dc:title>'
            f'<dc:identifier id="uid">urn:uuid:pdf-{uid}</dc:identifier>'
            '<dc:language>zh-CN</dc:language></metadata><manifest>'
            + "".join(manifest) + '</manifest><spine>'
            + "".join(spine) + '</spine></package>')
        nav = ('<html xmlns="http://www.w3.org/1999/xhtml"><body><nav '
               'epub:type="toc" xmlns:epub="http://www.idpf.org/2007/ops">'
               '<ol>' +
               "".join(
                   f'<li><a href="chapter-{i:04d}.xhtml">第 {i} 页</a></li>'
                   for i in range(1, len(chapters) + 1)) +
               '</ol></nav></body></html>')
        with zipfile.ZipFile(out, "w") as zf:
            info = zipfile.ZipInfo("mimetype")
            info.compress_type = zipfile.ZIP_STORED
            zf.writestr(info, "application/epub+zip")
            zf.writestr(
                "META-INF/container.xml",
                '<?xml version="1.0"?><container version="1.0" '
                'xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                '<rootfiles><rootfile full-path="OEBPS/content.opf" '
                'media-type="application/oebps-package+xml"/>'
                '</rootfiles></container>')
            zf.writestr("OEBPS/content.opf", opf)
            zf.writestr("OEBPS/nav.xhtml", nav)
            for i, chapter in enumerate(chapters, start=1):
                zf.writestr(f"OEBPS/chapter-{i:04d}.xhtml", chapter)
            for rel, data in assets.values():
                zf.writestr(f"OEBPS/{rel}", data)
    finally:
        doc.close()


def pdf_to_images(src: str, outdir: str, page_mode: str = "all",
                  page_number: int = 1, fmt: str = "PNG", zoom: int = 2,
                  progress: ProgressFn = _NOOP) -> int:
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf
    outdir = ensure_dir(outdir)
    doc = _open_pdf(src)
    stem = _stem(src)
    ext = fmt.lower()
    count = 0
    try:
        if page_mode == "single":
            indices = [page_number - 1]
            if indices[0] < 0 or indices[0] >= len(doc):
                raise AppError(f"只有 {len(doc)} 页，无法导出第 {page_number} 页")
        else:
            indices = list(range(len(doc)))
        matrix = pymupdf.Matrix(zoom, zoom)
        for idx in indices:
            pix = doc[idx].get_pixmap(matrix=matrix, alpha=False)
            dest = os.path.join(outdir, f"{stem}_第{idx + 1}页.{ext}")
            pix.save(dest)
            count += 1
            progress(f"已导出：{os.path.basename(dest)}")
    finally:
        doc.close()
    return count


def batch_pdf_to_dir(
    paths: list[str], outdir: str, kind: str,
    progress: ProgressFn = _NOOP,
) -> str:
    ext_map = {
        "Word": (".docx", pdf_to_word),
        "Excel": (".xlsx", pdf_to_excel),
        "PPT": (".pptx", pdf_to_ppt),
        "TXT": (".txt", pdf_to_txt),
        "HTML": (".html", pdf_to_html),
        "EPUB": (".epub", pdf_to_epub),
    }
    if kind not in ext_map:
        raise ValueError(f"不支持的转换类型：{kind}")
    ext, func = ext_map[kind]
    outdir = ensure_dir(outdir)
    for i, src in enumerate(paths, start=1):
        name = _stem(src)
        progress(f"（{i}/{len(paths)}）正在转换 {name} → {kind}…")
        func(src, os.path.join(outdir, name + ext), progress)
        progress(f"✓ 已生成：{name}{ext}")
    return f"已生成 {len(paths)} 个 {kind} 文件，输出位置：{outdir}"


def _docx():
    try:
        from docx import Document
    except ImportError as exc:
        raise AppError("缺少 python-docx，请运行 python bootstrap.py") from exc
    return Document


def _openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise AppError("缺少 openpyxl，请运行 python bootstrap.py") from exc
    return openpyxl
